"""
services/cash_flow_excel_service.py
Servicio de Exportación Excel para Flujo de Caja.
Genera un libro con:
- HOJA 1: Resumen de Flujo de Caja (Carry-Forward, Ingresos, Egresos, Saldo Final, Alerta de Déficit).
- HOJA 2: Detalle de Movimientos (Fecha, Naturaleza, Origen, Documento, Tercero, Cuenta, Descripción, Ingreso, Egreso, Estado, Conciliación).

Incluye:
- Estilo corporativo azul marino (#1E3A8A).
- Columnas numéricas nativas de Excel formateadas con formato de moneda.
- Sanitización estricta contra inyección de fórmulas (=, +, -, @, |, %).
"""

import io
from decimal import Decimal
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _safe_str(val: Any) -> str:
    """Previene inyección de fórmulas en Excel."""
    if val is None:
        return ""
    s = str(val).strip()
    if s.startswith(("=", "+", "-", "@", "|", "%")):
        return "'" + s
    return s


def export_cash_flow_to_excel(cf_data: Dict[str, Any], account_name: Optional[str] = None) -> bytes:
    wb = openpyxl.Workbook()

    # Estilos
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    meta_font = Font(name="Calibri", size=10, bold=True, color="1E293B")
    data_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    deficit_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    surplus_font = Font(name="Calibri", size=10, bold=True, color="16A34A")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    currency_format = "$#,##0"

    # ==========================================
    # HOJA 1: RESUMEN FLUJO DE CAJA
    # ==========================================
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Flujo de Caja"
    ws_resumen.views.sheetView[0].showGridLines = True

    # Encabezado
    filters = cf_data.get("filters", {})
    grouping_label = filters.get("grouping", "semanal").capitalize()
    acc_title = f"Cuenta: {account_name}" if account_name else "Cuenta: Todas las Cuentas (Consolidado)"
    
    ws_resumen.cell(row=1, column=1, value=_safe_str(f"REPORTE OFICIAL DE FLUJO DE CAJA ({grouping_label.upper()})")).font = Font(size=14, bold=True, color="1E3A8A")
    ws_resumen.cell(row=2, column=1, value=_safe_str(f"{acc_title} | Período: {filters.get('start_date')} al {filters.get('end_date')}")).font = meta_font

    kpis = cf_data.get("kpis", {})
    ws_resumen.cell(row=3, column=1, value=_safe_str(f"Saldo Inicial: ${Decimal(str(kpis.get('initial_balance', 0))):,.0f} | Saldo Proyectado Final: ${Decimal(str(kpis.get('final_projected_balance', 0))):,.0f} | Mínimo Proyectado: ${Decimal(str(kpis.get('min_projected_balance', 0))):,.0f}")).font = meta_font

    headers_summary = [
        "Período",
        "Saldo Inicial",
        "Ingresos Reales",
        "Ingresos Proyectados",
        "Total Ingresos",
        "Egresos Reales",
        "Egresos Proyectados",
        "Total Egresos",
        "Flujo Neto",
        "Saldo Final",
        "Estado / Alerta",
    ]

    for col_num, h in enumerate(headers_summary, 1):
        cell = ws_resumen.cell(row=5, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    row_idx = 6
    for p in cf_data.get("period_rows", []):
        c1 = ws_resumen.cell(row=row_idx, column=1, value=_safe_str(p.get("label")))
        c1.font = bold_font
        c1.border = thin_border

        cols_numeric = [
            (2, float(p.get("initial_balance", 0))),
            (3, float(p.get("real_inflows", 0))),
            (4, float(p.get("projected_inflows", 0))),
            (5, float(p.get("total_inflows", 0))),
            (6, float(p.get("real_outflows", 0))),
            (7, float(p.get("projected_outflows", 0))),
            (8, float(p.get("total_outflows", 0))),
            (9, float(p.get("net_period", 0))),
            (10, float(p.get("final_balance", 0))),
        ]
        for c_idx, val in cols_numeric:
            cell = ws_resumen.cell(row=row_idx, column=c_idx, value=val)
            cell.number_format = currency_format
            cell.font = data_font
            cell.alignment = Alignment(horizontal="right")
            cell.border = thin_border
            if c_idx == 10:
                cell.font = deficit_font if val < 0 else surplus_font

        # Estado / Alerta
        alert_str = "⚠ DÉFICIT DE CAJA" if p.get("has_deficit") else "OK"
        c11 = ws_resumen.cell(row=row_idx, column=11, value=alert_str)
        c11.font = deficit_font if p.get("has_deficit") else surplus_font
        c11.alignment = Alignment(horizontal="center")
        c11.border = thin_border

        row_idx += 1

    # Ajustar anchos
    for col in ws_resumen.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_resumen.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # ==========================================
    # HOJA 2: DETALLE DE MOVIMIENTOS
    # ==========================================
    ws_detalle = wb.create_sheet(title="Detalle Movimientos")
    ws_detalle.views.sheetView[0].showGridLines = True

    ws_detalle.cell(row=1, column=1, value=_safe_str("DETALLE DE MOVIMIENTOS FINANCIEROS")).font = Font(size=14, bold=True, color="1E3A8A")
    ws_detalle.cell(row=2, column=1, value=_safe_str(f"{acc_title} | Total Movimientos: {len(cf_data.get('movements', []))}")).font = meta_font

    headers_detail = [
        "Fecha",
        "Tipo",
        "Origen",
        "Documento",
        "Tercero / Entidad",
        "Cuenta Bancaria",
        "Descripción",
        "Ingreso ($)",
        "Egreso ($)",
        "Neto ($)",
        "Estado",
        "Conciliación",
    ]

    for col_num, h in enumerate(headers_detail, 1):
        cell = ws_detalle.cell(row=4, column=col_num, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    d_row_idx = 5
    for m in cf_data.get("movements", []):
        c_date = ws_detalle.cell(row=d_row_idx, column=1, value=_safe_str(m.get("date")))
        c_type = ws_detalle.cell(row=d_row_idx, column=2, value=_safe_str(m.get("flow_nature")))
        c_orig = ws_detalle.cell(row=d_row_idx, column=3, value=_safe_str(m.get("origin")))
        c_doc = ws_detalle.cell(row=d_row_idx, column=4, value=_safe_str(m.get("doc_number")))
        c_ent = ws_detalle.cell(row=d_row_idx, column=5, value=_safe_str(m.get("entity_name")))
        c_acc = ws_detalle.cell(row=d_row_idx, column=6, value=_safe_str(m.get("bank_label")))
        c_desc = ws_detalle.cell(row=d_row_idx, column=7, value=_safe_str(m.get("description")))

        for c in (c_date, c_type, c_orig, c_doc, c_ent, c_acc, c_desc):
            c.font = data_font
            c.border = thin_border

        inflow_val = float(m.get("inflow", 0))
        outflow_val = float(m.get("outflow", 0))
        net_val = float(m.get("net_amount", 0))

        c_in = ws_detalle.cell(row=d_row_idx, column=8, value=inflow_val if inflow_val > 0 else 0)
        c_out = ws_detalle.cell(row=d_row_idx, column=9, value=outflow_val if outflow_val > 0 else 0)
        c_net = ws_detalle.cell(row=d_row_idx, column=10, value=net_val)

        for c, val in ((c_in, inflow_val), (c_out, outflow_val), (c_net, net_val)):
            c.number_format = currency_format
            c.alignment = Alignment(horizontal="right")
            c.border = thin_border
            c.font = data_font

        c_st = ws_detalle.cell(row=d_row_idx, column=11, value=_safe_str(m.get("status")))
        c_rec = ws_detalle.cell(row=d_row_idx, column=12, value=_safe_str(m.get("reconciliation_status")))
        c_st.font = data_font
        c_st.border = thin_border
        c_st.alignment = Alignment(horizontal="center")
        c_rec.font = data_font
        c_rec.border = thin_border
        c_rec.alignment = Alignment(horizontal="center")

        d_row_idx += 1

    for col in ws_detalle.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_detalle.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
