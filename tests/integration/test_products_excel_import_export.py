import os
import io
import pytest
import openpyxl
from services.products_excel_service import (
    PRODUCT_EXCEL_COLUMNS,
    generate_products_excel,
    parse_and_validate_products_excel
)

def test_generate_products_template():
    """Verifica que la plantilla vacía contenga exactamente las columnas definidas y 2 filas de ejemplo."""
    excel_io = generate_products_excel([], is_template=True)
    wb = openpyxl.load_workbook(excel_io)
    ws = wb.active
    
    headers = [cell.value for cell in ws[1]]
    expected_headers = [col["header"] for col in PRODUCT_EXCEL_COLUMNS]
    assert headers == expected_headers
    
    # Debe tener 2 filas de ejemplo (filas 2 y 3)
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    sku_idx = headers.index("SKU")
    assert rows[1][sku_idx] == "MIE-500-001"
    assert rows[2][sku_idx] == "INS-ENV-001"

def test_generate_products_export():
    """Verifica la exportación de catálogo con datos reales y columnas de solo lectura."""
    sample_products = [
        {
            "id": 101,
            "sku": "PROD-TEST-1",
            "name": "Producto de Prueba 1",
            "category": "Mieles",
            "cost": 2500.0,
            "current_stock": 150.0,
            "ppp_cost": 2450.0,
            "unit_of_measure": "UN",
            "status": "Activo"
        }
    ]
    excel_io = generate_products_excel(sample_products, is_template=False)
    wb = openpyxl.load_workbook(excel_io)
    ws = wb.active
    
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 2  # Encabezado + 1 fila
    header = rows[0]
    data = rows[1]
    
    sku_idx = header.index("SKU")
    name_idx = header.index("Nombre del Producto")
    stock_idx = header.index("Stock Actual (Solo Lectura)")
    ppp_idx = header.index("Costo PPP (Solo Lectura)")
    
    assert data[sku_idx] == "PROD-TEST-1"
    assert data[name_idx] == "Producto de Prueba 1"
    assert data[stock_idx] == 150.0
    assert data[ppp_idx] == 2450.0

def test_parse_and_validate_products_excel_diff_and_protection():
    """Verifica detección de NUEVO, MODIFICADO, SIN_CAMBIOS, ERROR y protección de stock/ppp."""
    # Simular catálogo existente en BD
    by_sku = {
        "SKU-EXIST-1": {
            "id": 10,
            "sku": "SKU-EXIST-1",
            "name": "Miel 100g Original",
            "category": "Mieles",
            "cost": 1000.0,
            "current_stock": 50,
            "ppp_cost": 950.0,
            "min_stock": 10
        },
        "SKU-NO-CHANGE": {
            "id": 11,
            "sku": "SKU-NO-CHANGE",
            "name": "Miel 250g Sin Cambio",
            "category": "Mieles",
            "cost": 1500.0,
            "current_stock": 20,
            "ppp_cost": 1400.0,
            "min_stock": 5
        }
    }
    by_id = {
        10: by_sku["SKU-EXIST-1"],
        11: by_sku["SKU-NO-CHANGE"]
    }

    # Crear libro con varios casos:
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = [col["header"] for col in PRODUCT_EXCEL_COLUMNS]
    ws.append(headers)
    
    # Fila 1: Modificado
    row1 = [None] * len(headers)
    row1[headers.index("ID")] = 10
    row1[headers.index("SKU")] = "SKU-EXIST-1"
    row1[headers.index("Nombre del Producto")] = "Miel 100g Modificada"
    row1[headers.index("Categoría")] = "Mieles"
    row1[headers.index("Costo Unitario ($)")] = 1200.0  # cambió de 1000 a 1200
    row1[headers.index("Stock Actual (Solo Lectura)")] = 999999  # Intento malicioso de modificar stock
    row1[headers.index("Costo PPP (Solo Lectura)")] = 999999    # Intento malicioso de modificar PPP
    ws.append(row1)

    # Fila 2: Sin Cambios
    row2 = [None] * len(headers)
    row2[headers.index("ID")] = 11
    row2[headers.index("SKU")] = "SKU-NO-CHANGE"
    row2[headers.index("Nombre del Producto")] = "Miel 250g Sin Cambio"
    row2[headers.index("Categoría")] = "Mieles"
    row2[headers.index("Costo Unitario ($)")] = 1500.0
    ws.append(row2)

    # Fila 3: Nuevo
    row3 = [None] * len(headers)
    row3[headers.index("SKU")] = "SKU-NUEVO-99"
    row3[headers.index("Nombre del Producto")] = "Nuevo Producto Frasco"
    row3[headers.index("Categoría")] = "Envases"
    row3[headers.index("Tipo de Producto")] = "Insumo"
    row3[headers.index("Costo Unitario ($)")] = 350.0
    ws.append(row3)

    # Fila 4: Error (falta SKU y Nombre)
    row4 = [None] * len(headers)
    row4[headers.index("Categoría")] = "Huérfana"
    ws.append(row4)

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    result = parse_and_validate_products_excel(bio, by_sku, by_id)
    assert result["success"] is True
    summary = result["summary"]
    assert summary["total_rows"] == 4
    assert summary["nuevos"] == 1
    assert summary["modificados"] == 1
    assert summary["sin_cambios"] == 1
    assert summary["errores"] == 1

    items = result["items"]
    # Item 1: Modificado
    assert items[0]["action"] == "MODIFICADO"
    assert items[0]["target_id"] == 10
    changed_keys = [c["key"] for c in items[0]["changes"]]
    assert "name" in changed_keys
    assert "cost" in changed_keys
    # Asegurar que stock y ppp fueron omitidos y no aparecen en los cambios ni en el payload
    assert "current_stock" not in items[0]["payload"]
    assert "ppp_cost" not in items[0]["payload"]

    # Item 2: Sin Cambios
    assert items[1]["action"] == "SIN_CAMBIOS"

    # Item 3: Nuevo
    assert items[2]["action"] == "NUEVO"
    assert items[2]["sku"] == "SKU-NUEVO-99"

    # Item 4: Error
    assert items[3]["action"] == "ERROR"
    assert len(items[3]["errors"]) > 0

def test_routes_excel_export_and_preview(auth_client):
    """Verifica los endpoints web de exportación y plantilla en Flask."""
    # 1. Exportar completo
    res_export = auth_client.get('/productos/exportar')
    assert res_export.status_code == 200
    assert res_export.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in res_export.headers.get("Content-Disposition", "")

    # 2. Descargar plantilla
    res_template = auth_client.get('/productos/plantilla-excel')
    assert res_template.status_code == 200
    assert res_template.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "plantilla_productos.xlsx" in res_template.headers.get("Content-Disposition", "")
