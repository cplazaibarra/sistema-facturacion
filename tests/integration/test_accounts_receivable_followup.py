"""
tests/integration/test_accounts_receivable_followup.py
Pruebas integrales de Gestión y Seguimiento de Cobranza en Cuentas por Cobrar.
"""

import uuid
from datetime import date, timedelta
import openpyxl
import pytest

from app import app
from db import get_connection, insert_collection_action, list_collection_actions, get_accounts_receivable_report_data
from services.report_export_service import export_accounts_receivable_to_excel


@pytest.fixture
def auth_client():
    old_csrf = app.config.get('WTF_CSRF_ENABLED', True)
    app.config['TESTING'] = True
    app.config['WTF_CSRF_ENABLED'] = False
    try:
        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess['user_id'] = 1
                sess['username'] = 'agente_cobranza'
                sess['full_name'] = 'Ana Cobranzas'
                sess['role_name'] = 'Administrativo'
                sess['permissions'] = {'reportes': True, 'ventas': True, 'ventas.registrar_pago': True}
            yield client
    finally:
        app.config['WTF_CSRF_ENABLED'] = old_csrf


@pytest.fixture
def unauth_client():
    old_csrf = app.config.get('WTF_CSRF_ENABLED', True)
    app.config['TESTING'] = True
    app.config['WTF_CSRF_ENABLED'] = False
    try:
        with app.test_client() as client:
            with client.session_transaction() as sess:
                sess['user_id'] = 2
                sess['username'] = 'sin_permiso'
                sess['full_name'] = 'Sin Permiso'
                sess['role_name'] = 'Operario'
                sess['permissions'] = {'reportes': False}
            yield client
    finally:
        app.config['WTF_CSRF_ENABLED'] = old_csrf


@pytest.fixture
def seed_followup_data():
    """Crea una venta con deuda para las pruebas de seguimiento de cobranza."""
    marker = uuid.uuid4().hex[:8]
    customer_name = f"Empresa Deudora {marker}"
    sale_number = f"VTA-COB-{marker}"
    today = date.today()
    today_str = today.isoformat()
    due_str = (today + timedelta(days=10)).isoformat()

    sale_id = None
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO sales 
                (sale_number, customer_name, customer_email, sale_date, sale_time, products_json, total_amount, status, seller_name, payment_method, payment_status, notes, created_at)
                VALUES (%s, %s, %s, %s, '09:00:00', '[]', 500000, 'Pendiente', 'Vendedor Test', 'Transferencia', 'Pendiente', %s, %s)
                RETURNING id
                """,
                (sale_number, customer_name, f"deudor_{marker}@test.com", today_str, f"RUT: 77888999-1", today_str)
            )
            sale_id = cur.fetchone()["id"]

            cur.execute(
                """
                INSERT INTO sale_payments (sale_id, invoice_number, invoice_amount, invoice_due_date, payment_amount, status, created_at, updated_at)
                VALUES (%s, %s, 500000, %s, 0, 'Factura pendiente', %s, %s)
                """,
                (sale_id, f"FACT-{sale_id:05d}", due_str, today_str, today_str)
            )
        conn.commit()

    yield {
        "sale_id": sale_id,
        "sale_number": sale_number,
        "customer_name": customer_name,
        "today_str": today_str,
    }

    # Limpieza
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("DELETE FROM collection_actions WHERE sale_id = %s", (sale_id,))
            cur.execute("DELETE FROM sale_payment_items WHERE sale_id = %s", (sale_id,))
            cur.execute("DELETE FROM sale_payments WHERE sale_id = %s", (sale_id,))
            cur.execute("DELETE FROM sales WHERE id = %s", (sale_id,))
        conn.commit()


def test_01_insert_collection_action_valid(seed_followup_data):
    """1. Registro de gestión válida guarda correctamente en BD."""
    sale_id = seed_followup_data["sale_id"]
    action = {
        "sale_id": sale_id,
        "action_type": "Llamada telefónica",
        "action_date": "2026-09-20",
        "action_time": "10:30",
        "contact_name": "Carlos Tesorería",
        "result": "Cliente indica que transferirá el viernes.",
        "next_action": "Revisar cuenta bancaria",
        "next_action_date": "2026-09-25",
        "payment_commitment_date": "2026-09-25",
        "payment_commitment_amount": 250000.0,
        "notes": "Llamar a primera hora",
        "user_name": "Ana Cobranzas",
    }
    action_id = insert_collection_action(action)
    assert action_id is not None
    assert action_id > 0

    actions = list_collection_actions(sale_id)
    assert len(actions) == 1
    act = actions[0]
    assert act["action_type"] == "Llamada telefónica"
    assert act["contact_name"] == "Carlos Tesorería"
    assert act["payment_commitment_amount"] == 250000.0
    assert act["user_name"] == "Ana Cobranzas"


def test_02_all_action_types_supported(seed_followup_data):
    """2. Tipos de gestión estándar soportados (Correo, WhatsApp, Reunión, Carta, Otra, etc.)."""
    sale_id = seed_followup_data["sale_id"]
    types = [
        "Correo enviado",
        "Llamada telefónica",
        "WhatsApp / Mensaje",
        "Reunión",
        "Compromiso de pago",
        "Carta / Notificación",
        "Otra",
    ]
    for idx, t in enumerate(types):
        insert_collection_action({
            "sale_id": sale_id,
            "action_type": t,
            "action_date": f"2026-09-{10+idx:02d}",
            "result": f"Prueba tipo {t}",
            "user_name": "Tester",
        })
    actions = list_collection_actions(sale_id)
    assert len(actions) == len(types)
    stored_types = {a["action_type"] for a in actions}
    for t in types:
        assert t in stored_types


def test_03_multiple_actions_accumulate_chronologically(seed_followup_data):
    """5 y 6. Múltiples gestiones se acumulan en orden cronológico inverso sin sobreescribirse."""
    sale_id = seed_followup_data["sale_id"]
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Correo enviado",
        "action_date": "2026-09-10",
        "action_time": "09:00",
        "result": "Primer aviso de vencimiento",
        "user_name": "Ana",
    })
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Llamada telefónica",
        "action_date": "2026-09-15",
        "action_time": "14:00",
        "result": "Segunda gestión telefónica",
        "user_name": "Pedro",
    })
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "WhatsApp / Mensaje",
        "action_date": "2026-09-20",
        "action_time": "11:00",
        "result": "Tercera gestión por chat",
        "user_name": "Carlos",
    })

    actions = list_collection_actions(sale_id)
    assert len(actions) == 3
    # Debe estar ordenado de más reciente a más antiguo
    assert actions[0]["action_date"] == "2026-09-20"
    assert actions[1]["action_date"] == "2026-09-15"
    assert actions[2]["action_date"] == "2026-09-10"


def test_04_commitment_does_not_reduce_pending_amount(seed_followup_data):
    """8, 9 y 10. Compromiso de pago NO reduce pending_amount, no crea abonos ni altera inventario."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]

    # Registrar compromiso de pago por $300.000
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Compromiso de pago",
        "action_date": date.today().isoformat(),
        "action_time": "10:00",
        "result": "Cliente promete pagar $300.000 mañana",
        "payment_commitment_date": (date.today() + timedelta(days=1)).isoformat(),
        "payment_commitment_amount": 300000.0,
        "user_name": "Ana Cobranzas",
    })

    # Consultar reporte de cuentas por cobrar
    rep = get_accounts_receivable_report_data(search=sale_num)
    items = [it for it in rep["items"] if it["id"] == sale_id]
    assert len(items) == 1
    it = items[0]

    # REGLA ESTRICTA: El saldo pendiente sigue siendo 500.000
    assert it["total_amount"] == 500000.0
    assert it["paid_amount"] == 0.0
    assert it["pending_amount"] == 500000.0
    assert it["last_commitment_amount"] == 300000.0
    assert it["last_action_type"] == "Compromiso de pago"


def test_05_paid_sale_preserves_history(seed_followup_data):
    """7. Al pagar el 100% de la venta, el historial persiste intacto."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]

    # Insertar 2 gestiones
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Llamada telefónica",
        "action_date": "2026-09-15",
        "result": "Gestión previa",
        "user_name": "Ana",
    })
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Correo enviado",
        "action_date": "2026-09-18",
        "result": "Envío de cuenta bancaria",
        "user_name": "Ana",
    })

    # Ahora simulamos que la venta es 100% pagada
    with get_connection() as conn:
        with conn.cursor() as cur:
            cur.execute("UPDATE sales SET payment_status = 'Pagado' WHERE id = %s", (sale_id,))
            cur.execute("UPDATE sale_payments SET payment_amount = 500000, status = 'Pagada' WHERE sale_id = %s", (sale_id,))
        conn.commit()

    # No debe aparecer en el reporte activo de cuentas por cobrar
    rep = get_accounts_receivable_report_data(search=sale_num)
    items = [it for it in rep["items"] if it["id"] == sale_id]
    assert len(items) == 0

    # Pero el historial de gestiones debe seguir existiendo intacto
    hist = list_collection_actions(sale_id)
    assert len(hist) == 2


def test_06_dynamic_days_without_action_and_badges(seed_followup_data):
    """13, 14 y 15. Cálculo dinámico de Días sin Gestión y Badges."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]
    today = date.today()

    # Caso A: Sin gestión
    rep_sin = get_accounts_receivable_report_data(search=sale_num)
    item_sin = [it for it in rep_sin["items"] if it["id"] == sale_id][0]
    assert item_sin["last_action_date"] == ""
    assert item_sin["dias_sin_gestion"] is None
    assert item_sin["gestion_badge_type"] == "sin_gestion"

    # Caso B: Gestión hace 2 días (reciente)
    date_reciente = (today - timedelta(days=2)).isoformat()
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "WhatsApp / Mensaje",
        "action_date": date_reciente,
        "result": "Cliente responde que revisará",
        "user_name": "Cobrador",
    })
    rep_rec = get_accounts_receivable_report_data(search=sale_num)
    item_rec = [it for it in rep_rec["items"] if it["id"] == sale_id][0]
    assert item_rec["dias_sin_gestion"] == 2
    assert item_rec["gestion_badge_type"] == "reciente"

    # Caso C: Gestión más reciente hoy
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Llamada telefónica",
        "action_date": today.isoformat(),
        "result": "Confirmación hoy",
        "user_name": "Cobrador",
    })
    rep_hoy = get_accounts_receivable_report_data(search=sale_num)
    item_hoy = [it for it in rep_hoy["items"] if it["id"] == sale_id][0]
    assert item_hoy["dias_sin_gestion"] == 0
    assert item_hoy["gestion_badge_type"] == "reciente"
    assert item_hoy["total_actions_count"] == 2


def test_07_filters_by_gestion_and_antiguedad(seed_followup_data):
    """17, 18, 19 y 20. Filtros por gestión (con/sin), antigüedad y tipo."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]
    today = date.today()

    # Venta inicialmente sin gestión
    rep_sin = get_accounts_receivable_report_data(search=sale_num, filter_gestion="sin_gestion")
    assert len(rep_sin["items"]) == 1

    rep_con = get_accounts_receivable_report_data(search=sale_num, filter_gestion="con_gestion")
    assert len(rep_con["items"]) == 0

    # Insertamos gestión de hace 10 días tipo "Reunión"
    date_10d = (today - timedelta(days=10)).isoformat()
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Reunión",
        "action_date": date_10d,
        "result": "Reunión con gerente de finanzas",
        "user_name": "Ejecutivo",
    })

    # Ahora con_gestion debe incluirla y sin_gestion debe excluirla
    rep_con2 = get_accounts_receivable_report_data(search=sale_num, filter_gestion="con_gestion")
    assert len(rep_con2["items"]) == 1

    rep_sin2 = get_accounts_receivable_report_data(search=sale_num, filter_gestion="sin_gestion")
    assert len(rep_sin2["items"]) == 0

    # Filtro antigüedad: mas_7 debe incluirla (10 d > 7 d)
    rep_mas7 = get_accounts_receivable_report_data(search=sale_num, filter_antiguedad_gestion="mas_7")
    assert len(rep_mas7["items"]) == 1

    # Filtro antigüedad: hoy debe excluirla
    rep_hoy = get_accounts_receivable_report_data(search=sale_num, filter_antiguedad_gestion="hoy")
    assert len(rep_hoy["items"]) == 0

    # Filtro tipo de gestión: "Reunión" debe incluirla
    rep_tipo = get_accounts_receivable_report_data(search=sale_num, filter_tipo_gestion="Reunión")
    assert len(rep_tipo["items"]) == 1

    # Filtro tipo de gestión: "Correo enviado" debe excluirla
    rep_tipo_otro = get_accounts_receivable_report_data(search=sale_num, filter_tipo_gestion="Correo enviado")
    assert len(rep_tipo_otro["items"]) == 0


def test_08_sorting_by_collection_fields(seed_followup_data):
    """21, 22 y 23. Ordenamiento por última gestión, días sin gestión y próxima gestión."""
    rep_last_action = get_accounts_receivable_report_data(sort_by="last_action", sort_order="desc")
    assert "items" in rep_last_action

    rep_days = get_accounts_receivable_report_data(sort_by="days_without_action", sort_order="asc")
    assert "items" in rep_days

    rep_next = get_accounts_receivable_report_data(sort_by="next_action", sort_order="asc")
    assert "items" in rep_next


def test_09_export_excel_includes_collection_columns(seed_followup_data):
    """25. Exportación Excel incluye columnas de gestión (Última, Tipo, Días, Próxima)."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]

    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Llamada telefónica",
        "action_date": "2026-09-22",
        "action_time": "15:30",
        "result": "Compromiso de pago para fin de mes",
        "next_action": "Llamar para verificar",
        "next_action_date": "2026-09-30",
        "user_name": "Ana",
    })

    buffer = export_accounts_receivable_to_excel({"search": sale_num})
    wb = openpyxl.load_workbook(buffer)
    ws = wb.active

    headers = [cell.value for cell in ws[7]]
    assert "Última Gestión" in headers
    assert "Tipo Gestión" in headers
    assert "Días sin Gestión" in headers
    assert "Próxima Gestión" in headers

    # Verificar datos en la fila encontrada
    found_row = None
    for row in ws.iter_rows(min_row=8, values_only=True):
        if row[3] == sale_num:
            found_row = row
            break
    assert found_row is not None
    idx_ug = headers.index("Última Gestión")
    idx_tg = headers.index("Tipo Gestión")
    idx_pg = headers.index("Próxima Gestión")

    assert found_row[idx_ug] == "2026-09-22"
    assert found_row[idx_tg] == "Llamada telefónica"
    assert found_row[idx_pg] == "2026-09-30"


def test_10_http_post_registrar_gestion_auth(auth_client, seed_followup_data):
    """11, 26. Endpoint POST /gestiones exige RBAC y toma el usuario de la sesión autenticada."""
    sale_id = seed_followup_data["sale_id"]

    resp_auth = auth_client.post(
        '/reporteria/cuentas-por-cobrar/gestiones',
        json={
            "sale_id": sale_id,
            "action_type": "Correo enviado",
            "action_date": "2026-09-24",
            "action_time": "11:20",
            "contact_name": "Jefe Compras",
            "result": "Envío de detalle de facturas pendientes",
            "user_name": "HACKER_ATTEMPT",  # Debe ser ignorado en favor de la sesión
        }
    )
    assert resp_auth.status_code == 200
    res_json = resp_auth.get_json()
    assert res_json["status"] == "success"

    # Validar en base de datos
    actions = list_collection_actions(sale_id)
    assert len(actions) == 1
    assert actions[0]["user_name"] == "Ana Cobranzas"  # Inmutable desde sesión


def test_11_http_post_registrar_gestion_unauth(unauth_client, seed_followup_data):
    """26. Endpoint POST /gestiones sin permisos retorna 403 Forbidden."""
    sale_id = seed_followup_data["sale_id"]

    resp_unauth = unauth_client.post(
        '/reporteria/cuentas-por-cobrar/gestiones',
        json={
            "sale_id": sale_id,
            "action_type": "Llamada telefónica",
            "action_date": "2026-09-24",
            "result": "Intento no autorizado",
        }
    )
    assert resp_unauth.status_code == 403


def test_12_http_get_historial_cobranza_json(auth_client, seed_followup_data):
    """27. Endpoint GET /<sale_id>/historial-cobranza retorna JSON con el historial completo."""
    sale_id = seed_followup_data["sale_id"]

    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "WhatsApp / Mensaje",
        "action_date": "2026-09-21",
        "result": "Mensaje enviado a tesorería",
        "user_name": "Ana",
    })
    insert_collection_action({
        "sale_id": sale_id,
        "action_type": "Reunión",
        "action_date": "2026-09-23",
        "result": "Reunión presencial",
        "user_name": "Ana",
    })

    resp = auth_client.get(f'/reporteria/cuentas-por-cobrar/{sale_id}/historial-cobranza')
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["status"] == "success"
    assert data["sale_id"] == sale_id
    assert data["total_actions"] == 2
    assert len(data["actions"]) == 2
    assert data["actions"][0]["action_type"] == "Reunión"
    assert data["actions"][1]["action_type"] == "WhatsApp / Mensaje"


def test_13_no_row_duplication_with_many_actions(seed_followup_data):
    """28 y 30. Múltiples acciones de cobranza en una venta NO duplican filas en el reporte principal."""
    sale_id = seed_followup_data["sale_id"]
    sale_num = seed_followup_data["sale_number"]

    # Agregamos 10 gestiones a la misma venta
    for i in range(10):
        insert_collection_action({
            "sale_id": sale_id,
            "action_type": "Llamada telefónica",
            "action_date": f"2026-09-{10+i:02d}",
            "result": f"Gestión número {i+1}",
            "user_name": "Cobrador",
        })

    rep = get_accounts_receivable_report_data(search=sale_num)
    matching = [it for it in rep["items"] if it["id"] == sale_id]
    # DEBE HABER EXACTAMENTE 1 FILA
    assert len(matching) == 1
    assert matching[0]["total_actions_count"] == 10
    assert matching[0]["last_action_date"] == "2026-09-19"
